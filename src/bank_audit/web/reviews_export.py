"""Выгрузка жалоб вкладки «Отзывы» в Excel — в стиле AuditLens, с графиками.

Раньше выгрузка была CSV: таблица без оформления, без сводок и без указания,
откуда она. Теперь это рабочая книга:
  • «Обзор» — фирменная шапка, штамп «Экспортировано из AuditLens», фильтры
    выгрузки, плитки показателей и родные (редактируемые) графики Excel:
    главные проблемы, динамика по месяцам, площадки, города;
  • «Жалобы» — полный срез с разметкой ИИ: даты — датами, суммы — числами,
    эскалация подсвечена цветами интерфейса, фильтр и закреплённая шапка;
  • «Сводки» — таблицы под графиками с гистограммами в ячейках;
  • «О выгрузке» — фильтры, что значит каждый столбец, как считали.
Числа сводок считает код по тем же строкам, что лежат на листе «Жалобы».
"""
from __future__ import annotations

import io
from collections import Counter
from datetime import date

from .export_brand import MONO, SANS, SERIF, STAMP, XL, C, banner_png, stamp_line

SOURCE_LABEL = {"bankiru": "banki.ru", "banki_reviews": "banki.ru",
                "sravni_reviews": "sravni.ru", "bankiros_reviews": "bankiros.ru",
                "finuslugi_reviews": "finuslugi.ru"}

# Порядок и подписи столбцов листа «Жалобы» (ключи — из reviews_dash.export_rows).
COLUMNS = [
    ("дата", "Дата", 11), ("площадка", "Площадка", 12), ("город", "Город", 15),
    ("продукт", "Продукт", 18), ("главная проблема", "Главная проблема", 30),
    ("группа", "Группа проблем", 18), ("доп. проблемы", "Доп. проблемы", 24),
    ("эскалация", "Эскалация", 11), ("куда", "Куда обратился / грозит", 18),
    ("без согласия", "Без согласия", 10), ("ввели в заблуждение", "Ввели в заблуждение", 12),
    ("уязвимый клиент", "Уязвимый клиент", 15), ("сумма", "Сумма, ₽", 13),
    ("дата события", "Дата события", 12), ("суть", "Суть (пересказ ИИ)", 55),
    ("цитата", "Цитата клиента", 50), ("оценка", "Оценка", 8),
    ("вне кодификатора", "Вне кодификатора", 18), ("банк", "Банк", 14),
    ("ссылка", "Ссылка", 34), ("текст", "Полный текст отзыва", 80),
]
COLUMN_HELP = {
    "Дата": "дата публикации отзыва на площадке",
    "Площадка": "откуда отзыв: banki.ru, sravni.ru, finuslugi.ru и др.",
    "Главная проблема": "тема жалобы по кодификатору AuditLens (разметка ИИ)",
    "Доп. проблемы": "другие темы, названные в той же жалобе",
    "Эскалация": "«обратился» — клиент уже пожаловался в ЦБ, суд, надзор; «грозит» — грозит обратиться",
    "Куда обратился / грозит": "ЦБ, суд, Роспотребнадзор, ФАС, прокуратура, финомбудсмен, полиция",
    "Без согласия": "продукт или услуга подключены без согласия клиента",
    "Ввели в заблуждение": "клиент пишет, что условия объяснили неверно",
    "Уязвимый клиент": "пенсионер, инвалид, несовершеннолетний, участник СВО и др.",
    "Сумма, ₽": "сумма претензии, если клиент её назвал",
    "Суть (пересказ ИИ)": "краткий пересказ жалобы моделью",
    "Цитата клиента": "дословная фраза из отзыва, сверенная с текстом",
    "Вне кодификатора": "проблема, которой нет в кодификаторе тем",
}


def _date(s):
    try:
        return date.fromisoformat(str(s)[:10]) if s else None
    except ValueError:
        return None


def _month_label(ym: str) -> str:
    names = ("янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек")
    try:
        y, m = ym.split("-")
        return f"{names[int(m) - 1]} {y}"
    except (ValueError, IndexError):
        return ym


def describe_filters(f: dict) -> list[tuple[str, str]]:
    out = [("Банк", f.get("bank") or "Сбербанк"),
           ("Продукт", f.get("product") or "все продукты"),
           ("Тема", f.get("theme") or "все темы"),
           ("Признак", f.get("flag") or "—"),
           ("Площадка", SOURCE_LABEL.get(f.get("source") or "", f.get("source") or "все площадки")),
           ("Город", f.get("city") or "все города"),
           ("Период", (f"{f['days']} дней" if f.get("days") else
                       (_month_label(f["month"]) if f.get("month") else "всё время"))),
           ("Только эскалация", "да" if f.get("esc") else "нет")]
    return out


def _agg(rows: list[dict]) -> dict:
    themes = Counter(r.get("главная проблема") or "без темы" for r in rows)
    esc_by_theme = Counter(r.get("главная проблема") or "без темы" for r in rows
                           if r.get("эскалация"))
    months = Counter(str(r.get("дата") or "")[:7] for r in rows if r.get("дата"))
    sources = Counter(SOURCE_LABEL.get(r.get("площадка") or "", r.get("площадка") or "—")
                      for r in rows)
    cities = Counter(r.get("город") for r in rows if r.get("город"))
    products = Counter(r.get("продукт") for r in rows if r.get("продукт"))
    amounts = [float(r["сумма"]) for r in rows if isinstance(r.get("сумма"), (int, float))]
    dates = sorted(d for d in (_date(r.get("дата")) for r in rows) if d)
    return {
        "total": len(rows),
        "filed": sum(1 for r in rows if r.get("эскалация") == "обратился"),
        "threat": sum(1 for r in rows if r.get("эскалация") == "грозит"),
        "no_consent": sum(1 for r in rows if r.get("без согласия")),
        "misled": sum(1 for r in rows if r.get("ввели в заблуждение")),
        "vulnerable": sum(1 for r in rows if r.get("уязвимый клиент")),
        "amount_sum": sum(amounts), "amount_n": len(amounts),
        "themes": themes.most_common(15), "esc_by_theme": esc_by_theme,
        "months": sorted(months.items()), "sources": sources.most_common(8),
        "cities": cities.most_common(10), "products": products.most_common(10),
        "first": dates[0] if dates else None, "last": dates[-1] if dates else None,
    }


def _pct(a: int, b: int) -> float:
    return round(100.0 * a / b, 1) if b else 0.0


def to_xlsx(rows: list[dict], filters: dict, *, limit: int | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    a = _agg(rows)
    bank = filters.get("bank") or "Сбербанк"
    what = " · ".join(x for x in (filters.get("product"), filters.get("theme"),
                                   filters.get("flag")) if x)
    title = f"Жалобы клиентов: {bank}" + (f" · {what}" if what else "")
    period = ""
    if a["first"] and a["last"]:
        period = f"{a['first']:%d.%m.%Y} — {a['last']:%d.%m.%Y}"
    fl = describe_filters(filters)
    sub = " · ".join([f"{a['total']:,} жалоб".replace(",", " "), period,
                      next(v for k, v in fl if k == "Период"),
                      "разметка ИИ, проверенные цитаты"])
    wb = Workbook()
    wb.properties.creator = "AuditLens"
    wb.properties.lastModifiedBy = "AuditLens"
    wb.properties.title = title
    wb.properties.subject = "Жалобы клиентов — вкладка «Отзывы»"
    wb.properties.description = stamp_line()
    wb.properties.keywords = "AuditLens"

    # ── «Сводки»: таблицы-источники графиков ─────────────────────────────
    ss = wb.active
    ss.title = "Сводки"
    XL.sheet_base(ss, widths=[42, 12, 11, 14, 14, 3, 20, 12, 3, 22, 12], title=title)
    r = 1
    c = ss.cell(row=r, column=1, value="Сводки по выгрузке")
    c.font = XL.font(18, name=SERIF)
    XL.stamp(ss, 2, span=5)
    r = 4
    ranges: dict[str, tuple[int, int, int]] = {}      # имя → (строка шапки, n, столбец)

    def block(name: str, head: list[str], data: list[list], col: int, row: int,
              bar_col: int = 2, fmt: dict | None = None) -> int:
        e = ss.cell(row=row, column=col, value=name.upper())
        e.font = XL.font(8.5, color="accent", name=MONO)
        end = XL.table(ss, row + 1, head, data, col=col, formats=fmt or {})
        ss.auto_filter.ref = None
        if data:
            colL = get_column_letter(col + bar_col - 1)
            ss.conditional_formatting.add(
                f"{colL}{row + 2}:{colL}{row + 1 + len(data)}",
                DataBarRule(start_type="num", start_value=0, end_type="max",
                            color=C["accent"], showValue=True))
        ranges[name] = (row + 1, len(data), col)
        return end + 1

    th = [[t, n, _pct(n, a["total"]), a["esc_by_theme"].get(t, 0),
           _pct(a["esc_by_theme"].get(t, 0), n)] for t, n in a["themes"]]
    r = block("Главные проблемы", ["Проблема", "Жалоб", "Доля, %", "С эскалацией",
                                   "Эскалация, %"], th, 1, r,
              fmt={"Доля, %": "0.0", "Эскалация, %": "0.0"})
    mo = [[_month_label(m), n] for m, n in a["months"]]
    r2 = block("По месяцам", ["Месяц", "Жалоб"], mo, 7, 4)
    src = [[s, n] for s, n in a["sources"]]
    r2 = block("Площадки", ["Площадка", "Жалоб"], src, 7, r2)
    ci = [[s, n] for s, n in a["cities"]]
    r3 = block("Города", ["Город", "Жалоб"], ci, 10, 4)
    pr = [[s, n] for s, n in a["products"]]
    block("Продукты", ["Продукт", "Жалоб"], pr, 10, r3)

    # ── «Обзор» ──────────────────────────────────────────────────────────
    ov = wb.create_sheet("Обзор", 0)
    XL.sheet_base(ov, widths=[11.5] * 12, title=title)
    row = XL.banner(ov, banner_png(eyebrow="Отзывы · жалобы клиентов", title=title,
                                   subtitle=sub, width=1180), rows=8, width_px=1060)
    XL.stamp(ov, row, text_=stamp_line() + " · вкладка «Отзывы»", span=12)
    row += 1
    ftxt = " · ".join(f"{k.lower()}: {v}" for k, v in fl)
    fc = ov.cell(row=row, column=1, value=f"Фильтры: {ftxt}")
    fc.font = XL.font(9, color="ink3")
    fc.alignment = Alignment(wrap_text=True, vertical="top")
    ov.merge_cells(start_row=row, start_column=1, end_row=row, end_column=12)
    ov.row_dimensions[row].height = 28
    row += 2
    esc = a["filed"] + a["threat"]
    amount = (f"{a['amount_sum']:,.0f} ₽".replace(",", " ") if a["amount_n"] else "—")
    row = XL.kpis(ov, row, [
        ("Жалоб", a["total"], period or "в выгрузке"),
        ("С эскалацией", esc, (f"{_pct(esc, a['total'])}% · обратились {a['filed']}, "
                               f"грозят {a['threat']}")),
        ("Без согласия", a["no_consent"], f"{_pct(a['no_consent'], a['total'])}% жалоб"),
        ("Ввели в заблуждение", a["misled"], f"{_pct(a['misled'], a['total'])}% жалоб"),
        ("Уязвимые клиенты", a["vulnerable"], f"{_pct(a['vulnerable'], a['total'])}% жалоб"),
        ("Сумма претензий", amount, f"в {a['amount_n']} жалобах названа сумма"),
    ])
    from openpyxl.worksheet.pagebreak import Break
    ov.row_breaks.append(Break(id=row))          # графики — с новой страницы при печати
    row = XL.section(ov, row + 1, "Что болит у клиентов",
                     eyebrow="Графики · по данным листа «Сводки»", span=12)

    def bar(name: str, anchor: str, horizontal: bool = True, h_cm: float = 8.4,
            w_cm: float = 13.2, color: str | None = None) -> None:
        hdr, n, col = ranges[name]
        if not n:
            return
        ch = BarChart()
        ch.type = "bar" if horizontal else "col"
        ch.add_data(Reference(ss, min_col=col + 1, min_row=hdr, max_row=hdr + n),
                    titles_from_data=True)
        ch.set_categories(Reference(ss, min_col=col, min_row=hdr + 1, max_row=hdr + n))
        ch.gapWidth = 60
        ch.width, ch.height = w_cm, h_cm
        if horizontal:
            ch.x_axis.scaling.orientation = "maxMin"     # первая строка — сверху
        ch.y_axis.scaling.min = 0                        # шкала от нуля — без искажения
        ch.dataLabels = DataLabelList(showVal=True, showSerName=False, showCatName=False,
                                      showLegendKey=False, showPercent=False)
        XL.style_chart(ch, title=name, colors=[color] if color else None,
                       horizontal=horizontal)
        ov.add_chart(ch, anchor)

    bar("Главные проблемы", f"A{row}", h_cm=9.2)
    bar("По месяцам", f"G{row}", horizontal=False, h_cm=9.2, color=C["select"])
    row += 19
    ov.row_breaks.append(Break(id=row - 1))      # вторая пара графиков — своей страницей
    bar("Площадки", f"A{row}", color=C["ink2"])
    bar("Города", f"G{row}", color=C["sber"])
    row += 18
    note = ov.cell(row=row, column=1, value=(
        "Графики строятся по листу «Сводки» и пересчитываются, если изменить его. "
        "Жалобы — отзывы с оценкой 1–2★ и смешанные, размеченные моделью; тема — по "
        "кодификатору AuditLens, цитаты сверены с текстом отзыва."))
    note.font = XL.font(8.5, color="ink3")
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ov.merge_cells(start_row=row, start_column=1, end_row=row, end_column=12)
    ov.row_dimensions[row].height = 30

    # ── «Жалобы» ─────────────────────────────────────────────────────────
    ws = wb.create_sheet("Жалобы", 1)
    XL.sheet_base(ws, widths=[w for _, _, w in COLUMNS], title=title)
    t = ws.cell(row=1, column=1, value=title)
    t.font = XL.font(16, name=SERIF)
    ws.row_dimensions[1].height = 26
    XL.stamp(ws, 2, text_=f"{stamp_line()} · {a['total']} жалоб · {ftxt}", span=10)
    if limit and a["total"] >= limit:
        w_ = ws.cell(row=3, column=1, value=f"Выгружены последние {limit} жалоб — сузьте "
                                            f"фильтры, чтобы получить полный срез.")
        w_.font = XL.font(9, color="warn_ink", bold=True)
    head = [h for _, h, _ in COLUMNS]
    data = []
    for rr in rows:
        vals = []
        for k, _h, _w in COLUMNS:
            v = rr.get(k, "")
            if k in ("дата", "дата события"):
                v = _date(v) or (v or "")
            elif k == "площадка":
                v = SOURCE_LABEL.get(v or "", v or "")
            elif k in ("сумма", "оценка"):
                v = v if isinstance(v, (int, float)) else (v or "")
            vals.append(v)
        data.append(vals)
    hdr_row = 4
    end = XL.table(ws, hdr_row, head, data,
                   wrap={"Главная проблема", "Суть (пересказ ИИ)", "Цитата клиента",
                         "Доп. проблемы", "Куда обратился / грозит"},
                   formats={"Дата": "DD.MM.YYYY", "Дата события": "DD.MM.YYYY",
                            "Сумма, ₽": '#,##0 "₽"'},
                   links={"Ссылка"})
    ws.freeze_panes = ws.cell(row=hdr_row + 1, column=2)
    # Печать: без столбца полного текста (он на страницу не ляжет), шапка таблицы —
    # на каждой странице.
    ws.print_area = f"A1:{get_column_letter(len(head) - 1)}{max(end - 1, hdr_row + 1)}"
    ws.print_title_rows = f"{hdr_row}:{hdr_row}"
    if data:
        col = {h: get_column_letter(i) for i, h in enumerate(head, 1)}
        last = end - 1
        e = col["Эскалация"]
        rng = f"{e}{hdr_row + 1}:{e}{last}"
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="equal", formula=['"обратился"'], fill=XL.fill("accent_soft"),
            font=XL.font(9.5, bold=True, color="accent_ink")))
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="equal", formula=['"грозит"'], fill=XL.fill("warn_soft"),
            font=XL.font(9.5, bold=True, color="warn_ink")))
        for h in ("Без согласия", "Ввели в заблуждение"):
            L = col[h]
            ws.conditional_formatting.add(f"{L}{hdr_row + 1}:{L}{last}", CellIsRule(
                operator="equal", formula=['"да"'], font=XL.font(9.5, bold=True,
                                                                  color="accent_ink")))
        v = col["Уязвимый клиент"]
        ws.conditional_formatting.add(f"{v}{hdr_row + 1}:{v}{last}", FormulaRule(
            formula=[f'LEN({v}{hdr_row + 1})>0'], fill=XL.fill("warn_soft")))

    # ── «О выгрузке» ─────────────────────────────────────────────────────
    ab = wb.create_sheet("О выгрузке")
    XL.sheet_base(ab, widths=[26, 90], landscape=False, title=title)
    t = ab.cell(row=1, column=1, value="О выгрузке")
    t.font = XL.font(18, name=SERIF)
    XL.stamp(ab, 2, span=2)
    r = XL.section(ab, 4, "Фильтры", eyebrow="Что выгружено", span=2)
    r = XL.table(ab, r, ["Параметр", "Значение"], [list(x) for x in fl]) + 1
    ab.auto_filter.ref = None
    r = XL.section(ab, r, "Что значат столбцы", eyebrow="Лист «Жалобы»", span=2)
    r = XL.table(ab, r, ["Столбец", "Что это"],
                 [[k, v] for k, v in COLUMN_HELP.items()], wrap={"Что это"}) + 1
    ab.auto_filter.ref = None
    r = XL.section(ab, r, "Как считали", eyebrow="Методика", span=2)
    for line in (
        ("Жалоба — отзыв с оценкой 1–2★ или смешанный; разметку (тема, эскалация, признаки, "
         "пересказ, цитата) делает модель по кодификатору AuditLens, цитата сверяется с текстом."),
        ("Площадки: banki.ru, sravni.ru, bankiros.ru, finuslugi.ru — тот же корпус, что во "
         "вкладке «Отзывы»."),
        "Поиск по смыслу в выгрузку не входит: выгружается полный срез по фильтрам.",
        f"{STAMP}. Источник — вкладка «Отзывы».",
    ):
        cell = ab.cell(row=r, column=1, value=line)
        cell.font = XL.font(10, color="ink2")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ab.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ab.row_dimensions[r].height = 30
        r += 1

    # Книга открывается на «Обзоре»
    wb.active = 0
    for sh in wb.worksheets:
        sh.sheet_view.tabSelected = sh is ov
    ss.sheet_properties.tabColor = C["ink3"]
    ab.sheet_properties.tabColor = C["ink3"]
    ws.sheet_properties.tabColor = C["ink"]
    _ = (SANS,)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
