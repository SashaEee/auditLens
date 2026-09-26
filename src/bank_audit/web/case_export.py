"""Выгрузка аудит-дела в Excel и Word — в стиле AuditLens, с разметкой жалоб,
графиками и комментариями.

Excel — рабочая книга: «Дело» (шапка, показатели, графики, цель и разбор),
«Материалы» (одна строка на материал, фильтр, подсветка эскалаций), «Сводки»
(таблицы под графиками). Word — приложение к рабочему файлу проверки: обложка
как у PDF-отчёта, показатели, разбор дела, картина по материалам и карточки
материалов с сутью, цитатой и комментарием аудитора; фирменные шрифты
встроены в документ. На каждом листе и странице — «Экспортировано из
AuditLens». Собирается на сервере: у выгрузки один вид, кто бы её ни взял.
"""
from __future__ import annotations

import io
import re
from collections import Counter
from datetime import datetime, timezone

from .export_brand import MONO, SANS, SERIF, WD, XL, C, banner_png, bars_png, docx_bytes, stamp_line

_ESC = {"filed": "обратился", "threat": "грозит"}
_TO = {"cbr": "ЦБ", "court": "суд", "rpn": "Роспотребнадзор", "fas": "ФАС",
       "prosecutor": "прокуратура", "finombudsman": "финомбудсмен", "police": "полиция"}
_VULN = {"pensioner": "пенсионер", "low_income": "низкий доход", "svo": "участник СВО",
         "minor": "несовершеннолетний", "disabled": "инвалид", "ill": "тяжелобольной"}
_RISK = {"compliance": "комплаенс", "conduct": "практики", "ops": "операции"}
_KIND = {"document": "документ", "review": "жалоба", "offer": "продукт", "report": "отчёт"}


def _now() -> str:
    """Время выгрузки — по Москве, как штамп: в контейнере часы идут по UTC."""
    from ..clock import today_msk
    return today_msk().strftime("%d.%m.%Y %H:%M") + " МСК"


def _ru_date(d: str | None) -> str:
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", str(d or ""))
    return f"{m[3]}.{m[2]}.{m[1]}" if m else ""


def _amount(v) -> str:
    return f"{v:,.0f} ₽".replace(",", " ") if v is not None else ""


def _row(n: int, it: dict) -> dict:
    r = it.get("review") or {}
    doc = it["kind"] == "document"
    return {
        "№": n, "Тип": _KIND.get(it["kind"], it["kind"]),
        "Дата": _ru_date(r.get("date") or (str(it.get("fetched_at") or "")[:10] if doc else "")),
        "Банк": r.get("bank") or it.get("bank_name") or "",
        "Продукт": r.get("product") or "", "Город": r.get("city") or "",
        "Площадка": r.get("source") or "",
        "Главная проблема": r.get("issue_label") or "",
        "Класс риска": _RISK.get(r.get("risk") or "", ""),
        "Эскалация": _ESC.get(r.get("esc") or "", ""),
        "Куда": ", ".join(_TO.get(x, x) for x in r.get("esc_to") or []) if r.get("esc") in _ESC else "",
        "Уязвимый клиент": ", ".join(_VULN.get(x, x) for x in r.get("vulnerable") or []),
        "Без согласия": "да" if r.get("no_consent") else "",
        "Ввели в заблуждение": "да" if r.get("misled") else "",
        "Сумма": _amount(r.get("amount")),
        "Суть": r.get("summary") or it.get("title") or "",
        "Цитата": r.get("quote") or "",
        "Комментарий аудитора": it.get("note") or "",
        "Приобщил": it.get("added_by") or "",
        "Ссылка": it.get("url") or "",
    }


_WIDTH = {"№": 5, "Тип": 10, "Дата": 11, "Банк": 16, "Продукт": 18, "Город": 14,
          "Площадка": 12, "Главная проблема": 30, "Класс риска": 12, "Эскалация": 11,
          "Куда": 18, "Уязвимый клиент": 16, "Без согласия": 9, "Ввели в заблуждение": 11,
          "Сумма": 13, "Суть": 60, "Цитата": 50, "Комментарий аудитора": 40,
          "Приобщил": 14, "Ссылка": 40}


def _agg(case: dict, rows: list[dict]) -> dict:
    items = case.get("items") or []
    reviews = [it for it in items if it["kind"] == "review"]
    rv = [it.get("review") or {} for it in reviews]
    return {
        "items": len(items), "reviews": len(reviews),
        "other": len(items) - len(reviews),
        "esc": sum(1 for r in rv if r.get("esc") in _ESC),
        "filed": sum(1 for r in rv if r.get("esc") == "filed"),
        "vulnerable": sum(1 for r in rv if r.get("vulnerable")),
        "no_consent": sum(1 for r in rv if r.get("no_consent")),
        "amount": sum(float(r["amount"]) for r in rv if r.get("amount") is not None),
        "amount_n": sum(1 for r in rv if r.get("amount") is not None),
        "problems": Counter(r["Главная проблема"] for r in rows
                            if r["Главная проблема"]).most_common(12),
        "banks": Counter(r["Банк"] for r in rows if r["Банк"]).most_common(10),
        "types": Counter(r["Тип"] for r in rows).most_common(),
        "commented": sum(1 for r in rows if r["Комментарий аудитора"]),
    }


def _money(v: float) -> str:
    return f"{v:,.0f} ₽".replace(",", " ")


def _subtitle(case: dict, a: dict) -> str:
    return " · ".join(x for x in (
        f"владелец: {case.get('owner') or '—'}",
        f"материалов: {a['items']}",
        "открыто команде" if case.get("shared") else "личное дело",
        f"изменено {_ru_date(str(case.get('updated_at') or '')[:10])}"
        if case.get("updated_at") else "") if x)


def to_xlsx(case: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.pagebreak import Break

    rows = [_row(n, it) for n, it in enumerate(case.get("items") or [], 1)]
    a = _agg(case, rows)
    title = f"Аудит-дело: {case.get('title') or ''}"
    wb = Workbook()
    wb.properties.creator = "AuditLens"
    wb.properties.lastModifiedBy = "AuditLens"
    wb.properties.title = title
    wb.properties.subject = "Аудит-дело"
    wb.properties.description = stamp_line()
    wb.properties.keywords = "AuditLens"

    # ── «Сводки» — данные графиков ───────────────────────────────────────
    ss = wb.active
    ss.title = "Сводки"
    XL.sheet_base(ss, widths=[44, 12, 3, 30, 12, 3, 22, 12], title=title)
    ss.cell(row=1, column=1, value="Сводки по делу").font = XL.font(18, name=SERIF)
    XL.stamp(ss, 2, span=5)
    ranges: dict[str, tuple[int, int, int]] = {}

    def block(name: str, head: list[str], data: list[list], col: int, row: int) -> None:
        e = ss.cell(row=row, column=col, value=name.upper())
        e.font = XL.font(8.5, color="accent", name=MONO)
        XL.table(ss, row + 1, head, data, col=col)
        ss.auto_filter.ref = None
        ranges[name] = (row + 1, len(data), col)

    block("Материалы по проблемам", ["Главная проблема", "Материалов"],
          [list(x) for x in a["problems"]], 1, 4)
    block("По банкам", ["Банк", "Материалов"], [list(x) for x in a["banks"]], 4, 4)
    block("По типу", ["Тип", "Материалов"], [list(x) for x in a["types"]], 7, 4)

    # ── «Дело» ───────────────────────────────────────────────────────────
    ov = wb.create_sheet("Дело", 0)
    XL.sheet_base(ov, widths=[11.5] * 12, title=title)
    row = XL.banner(ov, banner_png(eyebrow="Аудит-дело", title=case.get("title") or "",
                                   subtitle=_subtitle(case, a), width=1180),
                    rows=8, width_px=1060)
    XL.stamp(ov, row, text_=stamp_line() + " · раздел «Аудит-дела»", span=12)
    row += 2
    row = XL.kpis(ov, row, [
        ("Материалов", a["items"], f"жалоб {a['reviews']} · прочих {a['other']}"),
        ("С эскалацией", a["esc"], f"уже обратились {a['filed']}"),
        ("Уязвимые клиенты", a["vulnerable"], "среди жалоб дела"),
        ("Без согласия", a["no_consent"], "продукт или услуга без согласия"),
        ("Сумма претензий", _money(a["amount"]) if a["amount_n"] else "—",
         f"названа в {a['amount_n']} жалобах"),
        ("С комментарием", a["commented"], "материалы с пометкой аудитора"),
    ])

    def text_block(r: int, eyebrow: str, head: str, body: str) -> int:
        r = XL.section(ov, r, head, eyebrow=eyebrow, span=12)
        for para in [x for x in re.split(r"\n\s*\n|\n", body or "") if x.strip()][:60]:
            plain = re.sub(r"\*\*(.+?)\*\*", r"\1", para).lstrip("#").strip()
            plain = re.sub(r"^[-*•]\s+", "•  ", plain)
            cell = XL.put(ov, r, 1, plain)
            is_h = para.lstrip().startswith("#")
            cell.font = XL.font(12 if is_h else 10, bold=is_h, color="ink" if is_h else "ink2",
                                name=SERIF if is_h else SANS)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ov.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
            ov.row_dimensions[r].height = max(16, 15 * (len(para) // 150 + 1))
            r += 1
        return r + 1

    if case.get("note"):
        row = text_block(row, "Цель", "Цель и примечание", case["note"])
    ov.row_breaks.append(Break(id=row - 1))
    row = XL.section(ov, row, "Картина по материалам", eyebrow="Графики · по листу «Сводки»",
                     span=12)

    def bar(name: str, anchor: str, color: str, w_cm: float = 13.2) -> None:
        hdr, n, col = ranges[name]
        if not n:
            return
        ch = BarChart()
        ch.type = "bar"
        ch.add_data(Reference(ss, min_col=col + 1, min_row=hdr, max_row=hdr + n),
                    titles_from_data=True)
        ch.set_categories(Reference(ss, min_col=col, min_row=hdr + 1, max_row=hdr + n))
        ch.gapWidth = 60
        ch.width, ch.height = w_cm, 8.4
        ch.x_axis.scaling.orientation = "maxMin"
        ch.y_axis.scaling.min = 0
        ch.dataLabels = DataLabelList(showVal=True, showSerName=False, showCatName=False,
                                      showLegendKey=False, showPercent=False)
        XL.style_chart(ch, title=name, colors=[color], horizontal=True)
        ov.add_chart(ch, anchor)

    bar("Материалы по проблемам", f"A{row}", C["accent"])
    bar("По банкам", f"G{row}", C["select"])
    row += 19
    if case.get("analysis"):
        ov.row_breaks.append(Break(id=row - 1))
        row = text_block(row, "Разбор ИИ", "Разбор материалов", case["analysis"])

    # ── «Материалы» ──────────────────────────────────────────────────────
    ws = wb.create_sheet("Материалы", 1)
    cols = list(_WIDTH)
    XL.sheet_base(ws, widths=[_WIDTH[k] for k in cols], title=title)
    ws.cell(row=1, column=1, value=title).font = XL.font(16, name=SERIF)
    ws.row_dimensions[1].height = 26
    XL.stamp(ws, 2, text_=f"{stamp_line()} · материалов: {len(rows)}", span=10)
    data = []
    for r in rows:
        vals = []
        for k in cols:
            v = r[k]
            if k == "Дата" and v:
                try:
                    v = datetime.strptime(v, "%d.%m.%Y").date()
                except ValueError:
                    pass
            vals.append(v)
        data.append(vals)
    hdr = 4
    end = XL.table(ws, hdr, cols, data,
                   wrap={"Главная проблема", "Суть", "Цитата", "Комментарий аудитора", "Куда"},
                   formats={"Дата": "DD.MM.YYYY"}, links={"Ссылка"})
    ws.freeze_panes = ws.cell(row=hdr + 1, column=2)
    ws.print_title_rows = f"{hdr}:{hdr}"
    if data:
        L = {k: get_column_letter(i) for i, k in enumerate(cols, 1)}
        last = end - 1
        e = f"{L['Эскалация']}{hdr + 1}:{L['Эскалация']}{last}"
        ws.conditional_formatting.add(e, CellIsRule(
            operator="equal", formula=['"обратился"'], fill=XL.fill("accent_soft"),
            font=XL.font(9.5, bold=True, color="accent_ink")))
        ws.conditional_formatting.add(e, CellIsRule(
            operator="equal", formula=['"грозит"'], fill=XL.fill("warn_soft"),
            font=XL.font(9.5, bold=True, color="warn_ink")))
        for k in ("Без согласия", "Ввели в заблуждение"):
            ws.conditional_formatting.add(f"{L[k]}{hdr + 1}:{L[k]}{last}", CellIsRule(
                operator="equal", formula=['"да"'],
                font=XL.font(9.5, bold=True, color="accent_ink")))
        v = L["Уязвимый клиент"]
        ws.conditional_formatting.add(f"{v}{hdr + 1}:{v}{last}", FormulaRule(
            formula=[f"LEN({v}{hdr + 1})>0"], fill=XL.fill("warn_soft")))
        cm = L["Комментарий аудитора"]
        ws.conditional_formatting.add(f"{cm}{hdr + 1}:{cm}{last}", FormulaRule(
            formula=[f"LEN({cm}{hdr + 1})>0"], fill=XL.fill("info_soft")))

    wb.active = 0
    for sh in wb.worksheets:
        sh.sheet_view.tabSelected = sh is ov
    ws.sheet_properties.tabColor = C["ink"]
    ss.sheet_properties.tabColor = C["ink3"]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _md_runs(par, text_: str, **kw) -> None:
    """**жирный** из разбора модели — жирным, остальное как есть."""
    for k, part in enumerate(re.split(r"\*\*(.+?)\*\*", text_)):
        if part:
            WD.run(par, part, bold=bool(k % 2), **kw)


def _material_card(doc, n: int, it: dict, r: dict) -> None:
    """Карточка материала: цветная кромка (эскалация — красная), рубрика,
    проблема, признаки, суть, цитата, комментарий аудитора, ссылка."""
    from docx.shared import Cm, Pt
    rv = it.get("review") or {}
    edge = ("accent" if rv.get("esc") == "filed" else
            "warn" if rv.get("esc") == "threat" or rv.get("vulnerable") else "hair2")
    t = doc.add_table(rows=1, cols=1)
    cell = t.rows[0].cells[0]
    cell.width = Cm(17)
    WD.shade(cell, "surface")
    WD.cell_borders(cell, left=(edge, 24), top=("hair", 4), bottom=("hair", 4),
                    right=("hair", 4))
    eyebrow = " · ".join(x for x in (f"[{n}] {r['Тип']}", r["Дата"], r["Банк"], r["Продукт"],
                                     r["Город"], r["Площадка"]) if x)
    p = cell.paragraphs[0]
    WD.run(p, eyebrow.upper(), font=MONO, size=7.5, color="ink3", spacing=0.4)
    head = r["Главная проблема"] or (it.get("title") or "")[:160]
    if head:
        hp = cell.add_paragraph()
        WD.run(hp, head, font=SANS, size=11, bold=True, color="ink")
        if r["Класс риска"]:
            WD.run(hp, f"  ·  {r['Класс риска']}", font=SANS, size=9, color="ink3")
    flags = [x for x in (
        (f"{r['Эскалация']}: {r['Куда']}" if r["Куда"] else r["Эскалация"]) if r["Эскалация"] else "",
        f"уязвимый клиент: {r['Уязвимый клиент']}" if r["Уязвимый клиент"] else "",
        "без согласия" if r["Без согласия"] else "",
        "ввели в заблуждение" if r["Ввели в заблуждение"] else "",
        f"сумма {r['Сумма']}" if r["Сумма"] else "") if x]
    if flags:
        fp = cell.add_paragraph()
        for i, f in enumerate(flags):
            WD.run(fp, ("  ·  " if i else "") + f, font=SANS, size=9, bold=True,
                   color="accent_ink" if ("обратился" in f or "согласия" in f
                                          or "заблуждение" in f) else "warn_ink")
    if r["Суть"] and r["Суть"] != head:
        sp = cell.add_paragraph()
        WD.run(sp, r["Суть"], font=SANS, size=10, color="ink2")
    if r["Цитата"]:
        qp = cell.add_paragraph()
        qp.paragraph_format.left_indent = Cm(0.4)
        from docx.oxml.ns import qn
        ppr = qp._element.get_or_add_pPr()
        bdr = ppr.makeelement(qn("w:pBdr"), {})
        bdr.append(bdr.makeelement(qn("w:left"), {qn("w:val"): "single", qn("w:sz"): "12",
                                                  qn("w:space"): "8",
                                                  qn("w:color"): C["hair2"]}))
        ppr.append(bdr)
        WD.run(qp, f"«{r['Цитата']}»", font=SERIF, size=10.5, color="ink", italic=False)
    if r["Комментарий аудитора"]:
        ct = cell.add_table(rows=1, cols=1)
        cc = ct.rows[0].cells[0]
        WD.shade(cc, "info_soft")
        WD.cell_borders(cc, left=("info", 12))
        WD.run(cc.paragraphs[0], "КОММЕНТАРИЙ АУДИТОРА", font=MONO, size=7, color="info",
               spacing=0.5)
        WD.run(cc.add_paragraph(), r["Комментарий аудитора"], font=SANS, size=10, color="ink")
        if r["Приобщил"]:
            WD.run(cc.add_paragraph(), f"приобщил: {r['Приобщил']}", font=MONO, size=7,
                   color="ink3")
    if r["Ссылка"]:
        lp = cell.add_paragraph()
        WD.run(lp, r["Ссылка"], font=MONO, size=7.5, color="select")
    for par in cell.paragraphs:
        par.paragraph_format.space_after = Pt(3)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def to_docx(case: dict) -> bytes:
    items = case.get("items") or []
    rows = [_row(n, it) for n, it in enumerate(items, 1)]
    a = _agg(case, rows)
    title = case.get("title") or "Аудит-дело"
    doc = WD.new(f"Аудит-дело: {title}")
    WD.cover(doc, eyebrow="Аудит-дело", title=title, ident=f"ДЕЛО № {case.get('case_id') or ''}",
             meta=[("Владелец", case.get("owner") or "—"),
                   ("Материалов", f"{a['items']} (жалоб {a['reviews']}, прочих {a['other']})"),
                   ("Доступ", "открыто команде" if case.get("shared") else "личное дело"),
                   ("Изменено", _ru_date(str(case.get("updated_at") or "")[:10]) or "—"),
                   ("Выгружено", _now())])
    WD.header_footer(doc, f"Аудит-дело «{title}»")

    WD.eyebrow(doc, "Коротко")
    h = doc.add_paragraph(style="Heading 1")
    WD.run(h, "Дело в цифрах", font=SERIF)
    WD.kpis(doc, [
        ("Материалов", str(a["items"]), f"жалоб {a['reviews']}"),
        ("С эскалацией", str(a["esc"]), f"обратились {a['filed']}"),
        ("Уязвимые", str(a["vulnerable"]), "клиенты"),
        ("Сумма претензий", _money(a["amount"]) if a["amount_n"] else "—",
         f"в {a['amount_n']} жалобах"),
    ])
    if case.get("note"):
        WD.eyebrow(doc, "Цель")
        p = doc.add_paragraph()
        _md_runs(p, case["note"], size=11, color="ink")
    if case.get("analysis"):
        WD.eyebrow(doc, "Разбор ИИ")
        h = doc.add_paragraph(style="Heading 1")
        WD.run(h, "Разбор материалов", font=SERIF)
        for line in case["analysis"].splitlines():
            s_ = line.strip()
            if not s_:
                continue
            if s_.startswith("#"):
                hp = doc.add_paragraph(style="Heading 2")
                WD.run(hp, s_.lstrip("#").strip(), font=SERIF)
            elif re.match(r"^[-*•]\s+", s_):
                _md_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*•]\s+", "", s_))
            else:
                _md_runs(doc.add_paragraph(), s_)
    if a["problems"] or a["banks"]:
        WD.eyebrow(doc, "Графики")
        h = doc.add_paragraph(style="Heading 1")
        WD.run(h, "Картина по материалам", font=SERIF)
        if a["problems"]:
            WD.picture(doc, bars_png("Материалы по главной проблеме", a["problems"],
                                     unit="", note="по разметке ИИ жалоб дела"))
        if len(a["banks"]) > 1:
            WD.picture(doc, bars_png("Материалы по банкам", a["banks"], color=C["select"]))
    doc.add_page_break()
    WD.eyebrow(doc, f"Материалы · {len(items)}")
    h = doc.add_paragraph(style="Heading 1")
    WD.run(h, "Материалы дела", font=SERIF)
    for n, (it, r) in enumerate(zip(items, rows), 1):
        _material_card(doc, n, it, r)
    end = doc.add_paragraph()
    WD.run(end, stamp_line(), font=MONO, size=8, color="accent")
    return docx_bytes(doc)
