"""«Пульс» → «Итоги»: пользуются ли AuditLens по направлениям аудита и
превращаются ли запросы в результат — то, что руководство спрашивает к
отчётам: предложения по доработкам и источникам, запросы, которые не могут
дать результата, разрез по направлениям, путь «вопрос → отчёт → аудит-дело».

Цифры считаются при открытии вкладки из живых таблиц: вопросы ИИ-помощнику,
чат и поиск «Аудита уязвимостей», аудит-дела, выгрузки и подписки «Аудита
отзывов», обращения, заявки на источники. Модель нужна в одном месте — понять
по тексту запроса направление и пригодность. Это разметка query_label: каждую
ночь (OUTCOMES_LABEL_HOUR_MSK) модель размечает новые запросы; владелец может
запустить её и вручную. Размеченное повторно не размечается, кроме названий
дел, которые переименовали.

Разметка в два прохода: быстрая модель — направления и пригодность всех
запросов, сильная перепроверяет только тех, кого быстрая сочла непригодными.
Одна быстрая модель завышала «плохие» вдвое: отчётность за прошлый год,
нормативку ЦБ и «да, хочу» в ответ на уточнение она записывала в непригодные.

Направление жалобы, которую аудитор положил в дело, берётся из темы дела
(«Доставка», «Дети 6–13 лет»): тема говорит, что аудитор проверяет, точнее
продукта самой жалобы. Только если тема ни к одному направлению не относится —
по продукту и срезу жалобы.
"""
from __future__ import annotations

import asyncio
import collections
import datetime as dt
import hashlib
import io
import logging
import os
import time

from sqlalchemy import text

from .. import db

log = logging.getLogger(__name__)
MSK = dt.timezone(dt.timedelta(hours=3))

# ── направления аудита ──────────────────────────────────────────────────────
DIRS: list[tuple[str, str]] = [
    ("mortgage", "Ипотека"), ("metals", "Операции с драгметаллами"),
    ("field", "Выездное обслуживание"), ("youth", "Дети и молодёжь"),
    ("selfemp", "Самозанятые"), ("credit_card", "Кредитные карты"),
    ("acquiring", "Эквайринг"), ("atm", "Устройства самообслуживания"),
    ("invest", "Инвестиционные продукты"), ("consumer", "Потребительские кредиты"),
    ("svo", "Участники СВО"), ("payroll", "Зарплатные проекты"),
    ("other", "Другое (вклады, дебетовые карты, переводы, РКО, страхование)"),
]
DIR_CODES = {c for c, _ in DIRS}

# продукт «Аудита отзывов» (подпись в фильтре) → направление
PRODUCT_DIR = {
    "Ипотека": "mortgage", "Драгметаллы и ОМС": "metals",
    "Выездной сервис — курьеры и выезд": "field", "Кредитная карта": "credit_card",
    "Бизнес: эквайринг": "acquiring", "УС — банкоматы и терминалы": "atm",
    "Инвестиции": "invest", "Потребительский кредит": "consumer", "Рассрочка": "consumer",
    "Зарплатные проекты": "payroll",
}
SLICE_DIR = {"payroll": "payroll", "us": "atm", "field": "field"}
VULNERABLE_DIR = {"minor": "youth", "svo": "svo"}

# пригодность запроса
QUALITY = {
    "ok": "Годный аудиторский запрос",
    "followup": "Уточнение к предыдущему ответу",
    "reference": "Простая справка без аудиторской задачи",
    "template": "Готовая карточка стартовой страницы",
    "vague": "Слишком общий — нет продукта, банка или цели",
    "internal": "Нужны закрытые данные банка",
    "personal": "Персональные данные — недопустимо",
    "offtopic": "Вне темы банковского аудита",
    "test": "Тест или случайный ввод",
}
BAD = ("vague", "internal", "personal", "offtopic", "test")
GOOD = ("ok", "followup")
# Карточки стартовой страницы ИИ-помощника до 25.09.2026: нажатие отправляло
# готовый вопрос. Две из четырёх сами были слишком общими — это не вопрос
# аудитора, поэтому в «плохие» такие нажатия не попадают.
OLD_CARDS = {
    "Сравни предложения по вкладам, выдели топ-5 и позицию Сбера.",
    "Какие основные жалобы у клиентов Сбербанка? Где подводные камни?",
    "Сравни ипотечные ставки между Сбером и рынком, выдели программы с господдержкой.",
    "Покажи изменения условий за последние 7 дней — что выросло, что упало.",
}
SECTIONS = {"ai": "ИИ-помощник", "vuln_chat": "Аудит уязвимостей", "vuln_search": "Аудит уязвимостей"}

LABEL_MODELS = [m.strip() for m in os.getenv(
    "OUTCOMES_LABEL_MODELS", "google/gemini-3-flash-preview,deepseek-ai/DeepSeek-V4-Pro").split(",") if m.strip()]
CHECK_MODELS = [m.strip() for m in os.getenv(
    "OUTCOMES_CHECK_MODELS", "openai/gpt-5.4,anthropic/claude-sonnet-4.6").split(",") if m.strip()]
LABEL_HOUR = int(os.getenv("OUTCOMES_LABEL_HOUR_MSK", "3"))
LABEL_MAX = int(os.getenv("OUTCOMES_LABEL_MAX", "1500"))       # запросов за прогон
_BATCH1, _BATCH2 = 20, 15

_DIR_RULES = (
    "mortgage — ипотека; metals — операции с драгметаллами (ОМС, слитки, монеты); "
    "field — выездное обслуживание (выездной менеджер, курьерская доставка карт и документов, выезд к клиенту); "
    "youth — дети и молодёжь (детские и подростковые карты, студенты, 14–24 года); "
    "selfemp — самозанятые (НПД); credit_card — кредитные карты; "
    "acquiring — эквайринг (торговый, интернет, QR и СБП для бизнеса); "
    "atm — устройства самообслуживания (банкоматы, терминалы, инфокиоски); "
    "invest — инвестиционные продукты (брокерский счёт, ИИС, ПИФ, облигации, ДУ, ИСЖ); "
    "consumer — потребительские кредиты (кредит наличными, POS, рассрочка); "
    "svo — участники СВО (военнослужащие, мобилизованные, их льготы); "
    "payroll — зарплатные проекты (зарплатные карты, зачисления от работодателя); "
    "other — другое (вклады, накопительные и дебетовые карты, переводы, РКО и МСБ, страхование, "
    "мошенничество без продукта, регулирование, рынок в целом, внутренние темы).")

_QUALITY_RULES = (
    "ok — понятны предмет и цель, ответ можно найти в открытых данных: условия и тарифы банков, отзывы, "
    "нормативка ЦБ и законы (в том числе правовые вопросы), опубликованная отчётность, новости, практики рынка;\n"
    "followup — ответ на уточняющий вопрос помощника или команда к прошлому ответу («да, хочу», «оставь 2026», "
    "«сделай таблицей», «подробнее»);\n"
    "vague — слишком общий: нет предмета, продукта, банка или цели, результат будет поверхностным;\n"
    "internal — нужны закрытые данные: внутренние регламенты (ВНД) и отчётность банка, транзакции и данные "
    "клиентов, внутренние метрики (NPS, CSI), доли и объёмы, которые не публикуются;\n"
    "personal — поиск сведений о конкретном частном лице (ФИО, дата рождения, связи) — недопустимо;\n"
    "reference — простой справочный факт без аудиторской задачи (ключевая ставка, «о чём эта программа»);\n"
    "offtopic — не о банках, финансовых услугах и аудите;\n"
    "test — тест, приветствие, бессмыслица, обрывок без смысла.")

_SYS1 = (
    "Ты размечаешь запросы аудиторов банка к аналитическому инструменту AuditLens: ИИ-помощник с поиском в "
    "интернете, корпус отзывов клиентов, тарифы банков, новости, база нормативных документов ЦБ, «аудит "
    "уязвимостей» — поиск схем обхода условий банков. Для каждого запроса верни:\n"
    "1) dirs — направления аудита, к которым относится запрос, 0–3 кода: " + _DIR_RULES + "\n"
    "2) quality — пригоден ли запрос, чтобы инструмент дал аудиторский результат:\n" + _QUALITY_RULES + "\n"
    "3) why — для quality кроме ok и followup: 3–8 слов, почему не даст результата.\n"
    "Строки вида «Название аудит-дела: …» — это тема проверки: для них важны только dirs, quality ставь ok.\n"
    'Ответ — строго JSON: {"items":[{"n":1,"dirs":["..."],"quality":"...","why":"..."}, ...]} по всем n по порядку.')

_SYS2 = (
    "Ты проверяешь, может ли аналитический инструмент для аудиторов банка (ИИ-помощник с поиском в интернете, "
    "корпус отзывов клиентов, тарифы банков, новости, база нормативных документов ЦБ) дать по запросу полезный "
    "аудиторский результат. Отчётность банков за прошлый год уже опубликована. Категории:\n" + _QUALITY_RULES +
    '\nwhy — 3–8 слов причины (для всех, кроме ok и followup).\n'
    'Ответ строго JSON: {"items":[{"n":1,"quality":"...","why":"..."}, ...]}.')


def _hash(s: str) -> str:
    return hashlib.md5(s.encode("utf-8")).hexdigest()


def _clean_dirs(v) -> list[str]:
    out: list[str] = []
    for d in v or []:
        if isinstance(d, str) and d in DIR_CODES and d not in out:
            out.append(d)
    return out[:3]


def merge_label(text_: str, first: dict | None, check: dict | None) -> dict | None:
    """Итог разметки запроса: направления — от быстрой модели, пригодность —
    от проверяющей, если она была; нажатие на старую карточку — «template»."""
    if not first:
        return None
    q = first.get("quality") if first.get("quality") in QUALITY else "ok"
    why, model = first.get("why"), first.get("model")
    if check and check.get("quality") in QUALITY:
        q, why, model = check["quality"], check.get("why"), check.get("model") or model
    if text_.strip() in OLD_CARDS:
        q, why = "template", None
    if q in GOOD:
        why = None
    return {"dirs": _clean_dirs(first.get("dirs")), "quality": q,
            "why": (str(why).strip()[:200] or None) if why else None, "model": model}


def query_dirs(dirs: list[str] | None, quality: str) -> list[str]:
    """Направления запроса для разреза: без явного — «Другое», кроме уточнений
    и тестов (у «да, хочу» направления нет, это продолжение прошлого запроса)."""
    ds = _clean_dirs(dirs)
    if ds:
        return ds
    return [] if quality in ("followup", "test") else ["other"]


def review_dirs(title_dirs: list[str] | None, product: str | None, slices: list[str] | None,
                vulnerable: list[str] | None) -> list[str]:
    """Направление жалобы в аудит-деле: по теме дела, иначе по самой жалобе."""
    td = [d for d in _clean_dirs(title_dirs) if d != "other"]
    if td:
        return td
    out = [SLICE_DIR[s] for s in (slices or []) if s in SLICE_DIR]
    if product in PRODUCT_DIR:
        out.append(PRODUCT_DIR[product])
    out += [VULNERABLE_DIR[v] for v in (vulnerable or []) if v in VULNERABLE_DIR]
    return _clean_dirs(out) or ["other"]


# ── разметка ────────────────────────────────────────────────────────────────
_PENDING_SQL = """
WITH src AS (
  SELECT 'ai' src, e.event_id ref_id, e.username, e.ts asked_at, e.payload->>'question' q
    FROM user_event e WHERE e.kind = 'ai_query' AND coalesce(e.payload->>'question', '') <> ''
  UNION ALL
  SELECT 'vuln_chat', m.message_id, w.user_id, m.created_at, m.content
    FROM loophole_chat_message m JOIN loophole_workspace w ON w.workspace_id = m.workspace_id
   WHERE m.role = 'user' AND coalesce(m.content, '') <> ''
  UNION ALL
  SELECT 'vuln_search', r.research_id, w.user_id, r.created_at, r.query_text
    FROM loophole_research r JOIN loophole_workspace w ON w.workspace_id = r.workspace_id
   WHERE coalesce(r.query_text, '') <> ''
  UNION ALL
  SELECT 'case_title', c.case_id, c.username, c.created_at, c.title
    FROM audit_case c WHERE c.deleted_at IS NULL AND coalesce(c.title, '') <> '')
SELECT s.* FROM src s
LEFT JOIN query_label l ON l.src = s.src AND l.ref_id = s.ref_id
WHERE l.ref_id IS NULL OR (s.src = 'case_title' AND l.text_hash <> md5(left(btrim(s.q, E' \t\r\n'), 2000)))
"""


def pending(limit: int | None = None) -> list[dict]:
    sql = _PENDING_SQL + " ORDER BY s.asked_at" + (" LIMIT :lim" if limit else "")
    with db.session() as s:
        return [dict(r) for r in s.execute(text(sql), {"lim": limit}).mappings()]


def pending_count() -> int:
    try:
        with db.session() as s:
            return int(s.execute(text(f"SELECT count(*) FROM ({_PENDING_SQL}) x")).scalar() or 0)
    except Exception:  # noqa: BLE001 — таблицы ещё нет (миграция не накатана)
        return 0


def _known(hashes: list[str]) -> dict[str, dict]:
    """Уже размеченные тексты: повтор того же вопроса не идёт в модель."""
    if not hashes:
        return {}
    with db.session() as s:
        rows = s.execute(text("""
            SELECT DISTINCT ON (text_hash) text_hash, dirs, quality, why, model
              FROM query_label WHERE text_hash = ANY(:h) AND src <> 'case_title'
             ORDER BY text_hash, labeled_at DESC"""), {"h": hashes}).mappings().all()
    return {r["text_hash"]: dict(r) for r in rows}


def _store(rows: list[dict]) -> None:
    with db.session() as s:
        for r in rows:
            s.execute(text("""
                INSERT INTO query_label (src, ref_id, username, asked_at, text_hash, question,
                                         dirs, quality, why, model)
                VALUES (:src, :ref_id, :username, :asked_at, :text_hash, :question,
                        :dirs, :quality, :why, :model)
                ON CONFLICT (src, ref_id) DO UPDATE SET text_hash = EXCLUDED.text_hash,
                    question = EXCLUDED.question, dirs = EXCLUDED.dirs, quality = EXCLUDED.quality,
                    why = EXCLUDED.why, model = EXCLUDED.model, labeled_at = now()"""), r)


class _Usage:
    def __init__(self):
        self.calls = self.tin = self.tout = 0


async def _ask(models: list[str], system: str, items: list[str], usage: _Usage,
               max_tokens: int) -> tuple[dict[int, dict], str | None]:
    """Пачка запросов → {номер: ответ}. Модели — по очереди, пока одна не
    ответит на 90% пачки."""
    from ..ai.llm_utils import _loose_json_loads
    from ..clock import today_anchor
    from ..digest.writer import _client
    user = "\n".join(f"[{k}] {t[:900]}" for k, t in enumerate(items, 1))
    for m in models:
        try:
            r = await _client().chat.completions.create(
                model=m, messages=[{"role": "system", "content": today_anchor() + "\n\n" + system},
                                   {"role": "user", "content": user}],
                temperature=0, max_tokens=max_tokens, response_format={"type": "json_object"})
            usage.calls += 1
            u = getattr(r, "usage", None)
            usage.tin += int(getattr(u, "prompt_tokens", 0) or 0)
            usage.tout += int(getattr(u, "completion_tokens", 0) or 0)
            j = _loose_json_loads(r.choices[0].message.content or "")
            got = {}
            for x in (j or {}).get("items") or []:
                try:
                    got[int(x["n"])] = x
                except (KeyError, TypeError, ValueError):
                    continue
            if len(got) >= len(items) * 0.9:
                return got, m
            log.warning("outcomes: %s ответила на %d из %d", m, len(got), len(items))
        except Exception as e:  # noqa: BLE001 — следующая модель
            log.warning("outcomes: %s не разметила пачку: %s", m, str(e)[:200])
    return {}, None


async def _label_texts(texts: list[str], usage: _Usage) -> dict[str, dict]:
    """Тексты → итоговая разметка (два прохода)."""
    sem = asyncio.Semaphore(4)
    first: dict[str, dict] = {}

    async def p1(batch: list[str]):
        async with sem:
            got, m = await _ask(LABEL_MODELS, _SYS1, batch, usage, 4000)
        for k, t in enumerate(batch, 1):
            if k in got:
                first[t] = {**got[k], "model": m}
    await asyncio.gather(*[p1(texts[i:i + _BATCH1]) for i in range(0, len(texts), _BATCH1)])
    doubt = [t for t, x in first.items() if x.get("quality") in BAD + ("reference",)
             and not t.startswith("Название аудит-дела:")]
    check: dict[str, dict] = {}

    async def p2(batch: list[str]):
        async with sem:
            got, m = await _ask(CHECK_MODELS, _SYS2, batch, usage, 6000)
        for k, t in enumerate(batch, 1):
            if k in got:
                check[t] = {**got[k], "model": m}
    await asyncio.gather(*[p2(doubt[i:i + _BATCH2]) for i in range(0, len(doubt), _BATCH2)])
    out = {}
    for t in texts:
        lab = merge_label(t, first.get(t), check.get(t))
        if lab and t.startswith("Название аудит-дела:"):
            lab.update(quality="ok", why=None)
        if lab:
            out[t] = lab
    return out


_RUN_LOCK = asyncio.Lock()


async def label_new(trigger: str = "night", limit: int | None = None) -> dict:
    """Разметить новые запросы. Повторный запуск во время идущего — пропуск."""
    if _RUN_LOCK.locked():
        return {"skipped": "уже идёт"}
    async with _RUN_LOCK:
        with db.session() as s:
            busy = s.execute(text("""SELECT 1 FROM query_label_run WHERE finished_at IS NULL
                                      AND started_at > now() - interval '30 minutes'""")).first()
            if busy:
                return {"skipped": "уже идёт в другом процессе"}
            run_id = s.execute(text("INSERT INTO query_label_run (trigger) VALUES (:t) RETURNING run_id"),
                               {"t": trigger}).scalar()
        usage, n, err = _Usage(), 0, None
        try:
            rows = await asyncio.to_thread(pending, limit or LABEL_MAX)
            for r in rows:
                q = (r["q"] or "").strip(" \t\r\n")[:2000]
                r["question"] = q
                r["text_hash"] = _hash(q)
                r["prompt"] = f"Название аудит-дела: {q}" if r["src"] == "case_title" else q
            known = await asyncio.to_thread(_known, sorted({r["text_hash"] for r in rows
                                                             if r["src"] != "case_title"}))
            todo = sorted({r["prompt"] for r in rows
                           if r["src"] == "case_title" or r["text_hash"] not in known})
            labels = await _label_texts(todo, usage) if todo else {}
            out = []
            for r in rows:
                lab = (labels.get(r["prompt"]) if r["src"] == "case_title" or r["text_hash"] not in known
                       else known[r["text_hash"]])
                if not lab:
                    continue
                out.append({"src": r["src"], "ref_id": r["ref_id"], "username": r["username"],
                            "asked_at": r["asked_at"], "text_hash": r["text_hash"],
                            "question": r["question"], "dirs": list(lab.get("dirs") or []),
                            "quality": lab["quality"], "why": lab.get("why"), "model": lab.get("model")})
            await asyncio.to_thread(_store, out)
            n = len(out)
        except Exception as e:  # noqa: BLE001
            err = f"{type(e).__name__}: {str(e)[:300]}"
            log.warning("outcomes: разметка не прошла: %s", err)
        with db.session() as s:
            s.execute(text("""UPDATE query_label_run SET finished_at = now(), labeled = :n, llm_calls = :c,
                                  tokens_in = :ti, tokens_out = :to, error = :e WHERE run_id = :id"""),
                      {"n": n, "c": usage.calls, "ti": usage.tin, "to": usage.tout, "e": err, "id": run_id})
        _CACHE.clear()
        log.info("outcomes: размечено %d, вызовов %d, токенов %d/%d", n, usage.calls, usage.tin, usage.tout)
        return {"labeled": n, "llm_calls": usage.calls, "error": err}


def last_run() -> dict | None:
    try:
        with db.session() as s:
            r = s.execute(text("""SELECT started_at, finished_at, trigger, labeled, error FROM query_label_run
                                   WHERE finished_at IS NOT NULL ORDER BY run_id DESC LIMIT 1""")).mappings().first()
        return {k: (v.isoformat() if isinstance(v, dt.datetime) else v) for k, v in r.items()} if r else None
    except Exception:  # noqa: BLE001
        return None


async def label_loop():
    """Каждую ночь в OUTCOMES_LABEL_HOUR_MSK:40 — новые запросы; после старта —
    догон, если последняя разметка старше суток."""
    await asyncio.sleep(600)
    while True:
        try:
            lr = await asyncio.to_thread(last_run)
            stale = not lr or dt.datetime.fromisoformat(lr["finished_at"]) < dt.datetime.now(MSK) - dt.timedelta(hours=24)
            if stale and await asyncio.to_thread(pending_count):
                await label_new("catch-up")
            now = dt.datetime.now(MSK)
            nxt = now.replace(hour=LABEL_HOUR, minute=40, second=0, microsecond=0)
            if nxt <= now:
                nxt += dt.timedelta(days=1)
            await asyncio.sleep((nxt - now).total_seconds())
            if await asyncio.to_thread(pending_count):
                await label_new("night")
        except asyncio.CancelledError:
            raise
        except Exception as e:  # noqa: BLE001
            log.warning("outcomes: цикл разметки: %s", e)
            await asyncio.sleep(1800)


# ── отчёт ───────────────────────────────────────────────────────────────────
def period(days: int, today: dt.date | None = None) -> tuple[dt.datetime | None, dt.datetime | None,
                                                            dt.datetime | None]:
    """(начало, конец, начало прошлого такого же периода) по МСК; days=0 — всё время."""
    if not days:
        return None, None, None
    today = today or dt.datetime.now(MSK).date()
    start = dt.datetime.combine(today - dt.timedelta(days=days - 1), dt.time(), MSK)
    return start, None, start - dt.timedelta(days=days)


_CACHE: dict[tuple, tuple[float, dict]] = {}


def _q(s, sql: str, p: dict) -> list[dict]:
    return [dict(r) for r in s.execute(text(sql), p).mappings()]


_PER = "{col} >= coalesce(CAST(:since AS timestamptz), '-infinity') AND {col} < coalesce(CAST(:until AS timestamptz), 'infinity')"


def _per(col: str) -> str:
    return _PER.format(col=col)


def _collect(ex: list[str], since, until) -> dict:
    """Сырьё отчёта за период — словари строк; считается в _build."""
    from ..rag import review_codebook
    schema = os.getenv("REVIEW_ANN_SCHEMA", review_codebook.VERSION)
    p = {"ex": ex or ["-"], "since": since, "until": until, "schema": schema}
    out: dict = {}
    with db.session() as s:
        out["labels"] = _q(s, f"""
            SELECT src, username, asked_at, question, dirs, quality, why FROM query_label
             WHERE src <> 'case_title' AND username <> ALL(:ex) AND {_per('asked_at')}
               AND (src = 'ai' OR username IN (SELECT username FROM app_user))""", p)
        out["titles"] = {r["ref_id"]: r["dirs"] for r in _q(s, "SELECT ref_id, dirs FROM query_label WHERE src = 'case_title'", p)}
        out["case_reviews"] = _q(s, f"""
            SELECT i.case_id, coalesce(i.added_by, c.username) username, ri.product, ri.slices,
                   (SELECT a.vulnerable FROM review_annotation a
                     WHERE a.url = i.url AND a.schema_version = :schema LIMIT 1) vulnerable
              FROM audit_case_item i JOIN audit_case c ON c.case_id = i.case_id
              LEFT JOIN review_index ri ON ri.url = i.url
             WHERE i.kind = 'review' AND c.deleted_at IS NULL
               AND coalesce(i.added_by, c.username) <> ALL(:ex) AND {_per('i.added_at')}""", p)
        out["review_ui"] = _q(s, f"""
            SELECT username, payload->>'action' act, payload->>'product' product,
                   (created_at AT TIME ZONE 'Europe/Moscow')::date d
              FROM usage_event
             WHERE kind = 'ui' AND page = 'reviews' AND username <> ALL(:ex)
               AND payload->>'action' IN ('reviews_export', 'reviews_product')
               AND coalesce(payload->>'product', '') <> '' AND {_per('created_at')}""", p)
        out["review_view_since"] = s.execute(text("""
            SELECT min(created_at) FROM usage_event WHERE kind = 'ui' AND page = 'reviews'
               AND payload->>'action' = 'reviews_product'""")).scalar()
        out["subs"] = _q(s, f"""SELECT username, product FROM review_subscription
                                 WHERE product <> '' AND username <> ALL(:ex) AND {_per('created_at')}""", p)
        out["cases"] = _q(s, f"""
            SELECT c.case_id, c.username, c.title, c.status, c.analysis_at IS NOT NULL analysed,
                   (SELECT count(*) FROM audit_case_item i WHERE i.case_id = c.case_id) items
              FROM audit_case c WHERE c.deleted_at IS NULL AND c.username <> ALL(:ex) AND {_per('c.created_at')}""", p)
        out["case_items"] = {r["kind"]: (r["n"], r["u"]) for r in _q(s, f"""
            SELECT i.kind, count(*) n, count(DISTINCT coalesce(i.added_by, c.username)) u
              FROM audit_case_item i JOIN audit_case c ON c.case_id = i.case_id
             WHERE c.deleted_at IS NULL AND coalesce(i.added_by, c.username) <> ALL(:ex) AND {_per('i.added_at')}
             GROUP BY 1""", p)}
        out["ai"] = _q(s, f"""
            SELECT count(*) n, count(DISTINCT username) u,
                   count(*) FILTER (WHERE payload->>'deep_requested' = 'true') deep
              FROM user_event WHERE kind = 'ai_query' AND username <> ALL(:ex) AND {_per('ts')}""", p)[0]
        out["runs"] = {r["st"]: r["n"] for r in _q(s, f"""
            SELECT payload->>'status' st, count(*) n FROM user_event
             WHERE kind = 'ai_run_end' AND username <> ALL(:ex) AND {_per('ts')} GROUP BY 1""", p)}
        out["vuln"] = _q(s, f"""
            SELECT count(*) n, count(DISTINCT u) u FROM (
              SELECT w.user_id u FROM loophole_chat_message m JOIN loophole_workspace w ON w.workspace_id = m.workspace_id
               WHERE m.role = 'user' AND {_per('m.created_at')}
              UNION ALL
              SELECT w.user_id FROM loophole_research r JOIN loophole_workspace w ON w.workspace_id = r.workspace_id
               WHERE {_per('r.created_at')}) x
             WHERE u <> ALL(:ex) AND u IN (SELECT username FROM app_user)""", p)[0]
        out["reports"] = _q(s, f"""SELECT count(*) n, count(DISTINCT username) u FROM report
                                    WHERE username <> ALL(:ex) AND {_per('created_at')}""", p)[0]
        out["reopened"] = s.execute(text(f"""
            SELECT count(DISTINCT payload->>'report_id') FROM user_event
             WHERE kind = 'report_open' AND username <> ALL(:ex) AND {_per('ts')}"""), p).scalar() or 0
        out["shares"] = _q(s, f"""SELECT count(*) n, count(DISTINCT report_id) r FROM report_share
                                   WHERE owner <> ALL(:ex) AND {_per('created_at')}""", p)[0]
        out["exports"] = _q(s, f"""
            SELECT count(*) n, count(DISTINCT username) u FROM usage_event
             WHERE kind = 'ui' AND page = 'reviews' AND payload->>'action' = 'reviews_export'
               AND username <> ALL(:ex) AND {_per('created_at')}""", p)[0]
        out["loophole"] = _q(s, f"""
            SELECT count(*) FILTER (WHERE action = 'mark_verdict') verdicts,
                   count(*) FILTER (WHERE action IN ('export_catalog_xlsx', 'research_report_export', 'export')) exports,
                   count(DISTINCT user_id) u
              FROM loophole_action_log
             WHERE user_id <> ALL(:ex) AND user_id IN (SELECT username FROM app_user) AND {_per('created_at')}""", p)[0]
        out["ratings"] = {f"{r['kind']}/{r['verdict']}": r["n"] for r in _q(s, f"""
            SELECT kind, verdict, count(*) n FROM item_feedback
             WHERE kind IN ('ai_answer', 'news') AND username <> ALL(:ex) AND {_per('created_at')}
             GROUP BY 1, 2""", p)}
        out["tickets"] = _q(s, """
            SELECT ticket_id, created_at, kind, section_label, body, status FROM user_ticket
             WHERE username <> ALL(:ex) ORDER BY ticket_id DESC""", p)
        out["proposals"] = _q(s, """
            SELECT proposal_id, created_at, purpose, domain, title, status, review_note FROM source_proposal
             WHERE proposed_by <> ALL(:ex) ORDER BY proposal_id DESC""", p)
    return out


TICKET_ST = {"new": "ждёт ответа", "accepted": "принято", "in_progress": "в работе", "done": "сделано",
             "wontfix": "не будем делать", "exists": "уже есть"}
PROPOSAL_ST = {"pending": "ждёт решения", "approved": "одобрен", "rejected": "отклонён"}
PURPOSE = {"reviews": "отзывы", "tariffs": "тарифы", "digest": "новостные обзоры", "ai": "ИИ-помощник"}


def _build(raw: dict, since=None) -> dict:
    """Сырьё → разделы отчёта. Без БД — тестируется на подставных строках."""
    def in_per(ts) -> bool:
        return since is None or (ts is not None and ts >= since)
    cols = {k: collections.Counter() for k in ("ai", "vuln", "rev", "cases")}
    users = {k: collections.defaultdict(set) for k in cols}
    parts: dict[str, collections.Counter] = collections.defaultdict(collections.Counter)
    qual = {"ai": collections.Counter(), "vuln": collections.Counter()}
    bad: dict[tuple, dict] = {}
    for r in raw["labels"]:
        k = "ai" if r["src"] == "ai" else "vuln"
        qual[k][r["quality"]] += 1
        for d in query_dirs(r["dirs"], r["quality"]):
            cols[k][d] += 1
            users[k][d].add(r["username"])
        if r["quality"] in BAD:
            key = (k, r["question"].strip().lower())
            b = bad.get(key)
            if b:
                b["n"] += 1
                b["d"] = max(b["d"], r["asked_at"])
            else:
                bad[key] = {"d": r["asked_at"], "section": SECTIONS[r["src"]], "quality": r["quality"],
                            "text": r["question"][:400], "why": r["why"], "n": 1}
    for r in raw["case_reviews"]:
        for d in review_dirs(raw["titles"].get(r["case_id"]), r["product"], r["slices"], r["vulnerable"]):
            cols["rev"][d] += 1
            users["rev"][d].add(r["username"])
            parts[d]["case"] += 1
    seen_view = set()
    for r in raw["review_ui"]:
        d = PRODUCT_DIR.get(r["product"], "other")
        if r["act"] == "reviews_product":
            key = (r["username"], r["product"], r["d"])        # день — один просмотр
            if key in seen_view:
                continue
            seen_view.add(key)
            parts[d]["view"] += 1
        else:
            parts[d]["export"] += 1
        cols["rev"][d] += 1
        users["rev"][d].add(r["username"])
    for r in raw["subs"]:
        d = PRODUCT_DIR.get(r["product"], "other")
        cols["rev"][d] += 1
        users["rev"][d].add(r["username"])
        parts[d]["sub"] += 1
    for c in raw["cases"]:
        for d in query_dirs(raw["titles"].get(c["case_id"]), "ok"):
            cols["cases"][d] += 1
            users["cases"][d].add(c["username"])
    dirs = [{"code": c, "name": n,
             **{k: cols[k][c] for k in cols}, **{k + "_users": len(users[k][c]) for k in cols},
             "rev_parts": dict(parts[c])} for c, n in DIRS]

    def qsum(c: collections.Counter) -> dict:
        total = sum(c.values())
        b = sum(c[x] for x in BAD)
        return {"total": total, "counts": {k: c[k] for k in QUALITY}, "bad": b,
                "good": c["ok"] + c["followup"],
                "good_share": round((c["ok"] + c["followup"]) / total, 4) if total else None,
                "bad_share": round(b / total, 4) if total else None}
    cases = raw["cases"]
    st = collections.Counter(c["status"] for c in cases)
    ci = raw["case_items"]
    runs = raw["runs"]
    tickets = raw["tickets"]
    tst = collections.Counter(t["status"] for t in tickets)
    props = raw["proposals"]
    pst = collections.Counter(p["status"] for p in props)
    return {
        "dirs": dirs,
        "review_view_since": raw["review_view_since"].isoformat() if raw.get("review_view_since") else None,
        "quality": {"ai": qsum(qual["ai"]), "vuln": qsum(qual["vuln"])},
        "bad_examples": [{**b, "d": b["d"].isoformat()} for b in
                         sorted(bad.values(), key=lambda b: b["d"], reverse=True)],
        "funnel": {
            "questions": int(raw["ai"]["n"]), "askers": int(raw["ai"]["u"]), "deep": int(raw["ai"]["deep"]),
            "runs_ok": int(runs.get("ok", 0)), "runs_failed": int(runs.get("failed", 0)) + int(runs.get("error", 0)),
            "runs_stopped": int(runs.get("stopped", 0)),
            "reports": int(raw["reports"]["n"]), "report_users": int(raw["reports"]["u"]),
            "reopened": int(raw["reopened"]), "shared": int(raw["shares"]["r"]), "shares": int(raw["shares"]["n"]),
            "case_reports": int(ci.get("report", (0, 0))[0]), "case_reviews": int(ci.get("review", (0, 0))[0]),
            "case_documents": int(ci.get("document", (0, 0))[0]),
            "cases": len(cases), "case_owners": len({c["username"] for c in cases}),
            "cases_collect": st.get("collect", 0), "cases_work": st.get("work", 0),
            "cases_done": sum(n for k, n in st.items() if k not in ("collect", "work")),
            "cases_analysed": sum(1 for c in cases if c["analysed"]),
            "cases_empty": sum(1 for c in cases if not c["items"]),
            "exports": int(raw["exports"]["n"]), "export_users": int(raw["exports"]["u"]),
            "subs": len(raw["subs"]), "sub_users": len({s["username"] for s in raw["subs"]}),
            "vuln_queries": int(raw["vuln"]["n"]), "vuln_users": int(raw["vuln"]["u"]),
            "verdicts": int(raw["loophole"]["verdicts"] or 0), "vuln_exports": int(raw["loophole"]["exports"] or 0),
            "loophole_users": int(raw["loophole"]["u"] or 0),
        },
        "case_topics": [{"title": c["title"], "items": int(c["items"]), "status": c["status"]}
                        for c in sorted(cases, key=lambda c: -int(c["items"]))[:12] if c["items"]],
        "ratings": raw["ratings"],
        "tickets": {"total": len(tickets), "done": tst.get("done", 0), "new_in_period":
                    sum(1 for t in tickets if in_per(t["created_at"])),
                    "counts": {TICKET_ST.get(k, k): n for k, n in tst.items()},
                    "items": [{"id": t["ticket_id"], "d": t["created_at"].isoformat(), "section": t["section_label"],
                               "body": (t["body"] or "").strip()[:300], "status": TICKET_ST.get(t["status"], t["status"]),
                               "done": t["status"] in ("done", "exists")} for t in tickets]},
        "sources": {"total": len(props), "approved": pst.get("approved", 0), "pending": pst.get("pending", 0),
                    "rejected": pst.get("rejected", 0),
                    "new_in_period": sum(1 for x in props if in_per(x["created_at"])),
                    "items": [{"id": x["proposal_id"], "d": x["created_at"].isoformat(), "domain": x["domain"],
                               "title": x["title"], "purpose": PURPOSE.get(x["purpose"], x["purpose"]),
                               "status": PROPOSAL_ST.get(x["status"], x["status"]), "note": x["review_note"]}
                              for x in props if x["status"] != "rejected"]},
    }


def report(viewer: str | None = None, with_me: bool | None = None, days: int = 0) -> dict:
    from . import telemetry
    days = days if days in (0, 7, 30, 90) else 0
    ex = telemetry.excluded(viewer, with_me)
    key = (days, tuple(ex))
    hit = _CACHE.get(key)
    if hit and time.time() - hit[0] < 120:
        return hit[1]
    since, until, prev = period(days)
    res = _build(_collect(ex, since, until), since)
    res["period"] = {"days": days, "since": since.isoformat() if since else None}
    if days:
        pr = _build(_collect(ex, prev, since), prev)
        res["prev"] = {"questions": pr["funnel"]["questions"], "askers": pr["funnel"]["askers"],
                       "vuln_queries": pr["funnel"]["vuln_queries"], "cases": pr["funnel"]["cases"],
                       "case_items": pr["funnel"]["case_reports"] + pr["funnel"]["case_reviews"]
                       + pr["funnel"]["case_documents"],
                       "good_share": pr["quality"]["ai"]["good_share"]}
    res["labels"] = {"last_run": last_run(), "pending": pending_count(), "hour_msk": LABEL_HOUR}
    _CACHE[key] = (time.time(), res)
    return res


# ── выгрузка в Excel ────────────────────────────────────────────────────────
def to_xlsx(rep: dict) -> bytes:
    from openpyxl import Workbook

    from .export_brand import XL, banner_png, stamp_line
    days = rep["period"]["days"]
    per = "за всё время" if not days else f"за {days} дн"
    title = f"AuditLens: итоги использования {per}"
    f = rep["funnel"]
    wb = Workbook()
    wb.properties.creator = "AuditLens"
    wb.properties.title = title
    ws = wb.active
    ws.title = "Итоги"
    XL.sheet_base(ws, widths=[40, 14, 14, 14, 14, 14, 14, 14, 14], title=title)
    row = XL.banner(ws, banner_png(eyebrow="Пульс · итоги", title="Итоги использования AuditLens",
                                   subtitle=per), width_px=1100)
    XL.stamp(ws, row, text_=stamp_line(), span=9)
    qa = rep["quality"]["ai"]
    row = XL.kpis(ws, row + 2, [
        ("Вопросы ИИ-помощнику", f["questions"], f"{f['askers']} чел."),
        ("Запросы в «Аудите уязвимостей»", f["vuln_queries"], f"{f['vuln_users']} чел."),
        ("Годные вопросы ИИ", f"{round((qa['good_share'] or 0) * 100)}%", "годные и уточнения"),
        ("Аудит-дела", f["cases"], f"{f['case_owners']} аудиторов"),
    ], span_each=2)
    row = XL.section(ws, row, "Запросы по направлениям", span=9)
    data = []
    for d in rep["dirs"]:
        p = d["rev_parts"]
        data.append([d["name"], d["ai"], d["ai_users"], d["rev"], d["rev_users"],
                     p.get("case", 0), p.get("view", 0), d["vuln"], d["vuln_users"]])
    hdr = row
    row = XL.table(ws, row, ["Направление", "ИИ-помощник", "чел.", "Аудит отзывов", "чел.", "из них жалоб в делах",
                             "из них просмотров продукта", "Аудит уязвимостей", "чел."], data, wrap={"Направление"})
    ws.row_dimensions[hdr].height = 44
    note = ("Аудит отзывов — жалобы в аудит-делах (направление — по теме дела), выгрузки в Excel и подписки "
            "с продуктом" + (", просмотры продукта в фильтре" if rep.get("review_view_since") else "") +
            ". Один запрос может относиться к нескольким направлениям.")
    c = ws.cell(row=row, column=1, value=note)
    c.font = XL.font(9, color="ink3", italic=True)
    row = XL.section(ws, row + 2, "От запроса к результату", span=9)
    fr = [["Вопросы ИИ-помощнику", f["questions"], f"{f['askers']} чел.; заказов глубокого отчёта — {f['deep']}"],
          ["Ответ получен", f["runs_ok"], f"остановлено {f['runs_stopped']}, сбоев {f['runs_failed']}"],
          ["Сохранено отчётов", f["reports"], f"{f['report_users']} чел."],
          ["Открыты повторно", f["reopened"], ""],
          ["Пересланы коллегам", f["shared"], f"{f['shares']} пересылок"],
          ["Отчёты в аудит-делах", f["case_reports"], ""],
          ["Жалобы клиентов в аудит-делах", f["case_reviews"], ""],
          ["Документы в аудит-делах", f["case_documents"], ""],
          ["Аудит-дела", f["cases"], f"сбор материалов {f['cases_collect']}, в работе {f['cases_work']}, "
                                     f"с итоговым анализом {f['cases_analysed']}, пустых {f['cases_empty']}"],
          ["Выгрузки отзывов в Excel", f["exports"], f"{f['export_users']} чел."],
          ["Подписки на сигналы по отзывам", f["subs"], f"{f['sub_users']} чел."],
          ["Отметки экспертов по схемам («Аудит уязвимостей»)", f["verdicts"], ""]]
    XL.table(ws, row, ["Шаг", "Сколько", "Комментарий"], fr, wrap={"Комментарий"})
    # качество
    wq = wb.create_sheet("Качество запросов")
    XL.sheet_base(wq, widths=[52, 16, 10, 18, 10], title=title)
    r = XL.section(wq, 1, "Качество запросов", span=5)
    qv = rep["quality"]["vuln"]
    rows = [[QUALITY[k], qa["counts"][k], f"{round(100 * qa['counts'][k] / qa['total'])}%" if qa["total"] else "",
             qv["counts"][k], f"{round(100 * qv['counts'][k] / qv['total'])}%" if qv["total"] else ""] for k in QUALITY]
    rows.append(["Не дают результата", qa["bad"], f"{round(100 * (qa['bad_share'] or 0))}%",
                 qv["bad"], f"{round(100 * (qv['bad_share'] or 0))}%"])
    rows.append(["Всего размечено", qa["total"], "", qv["total"], ""])
    XL.table(wq, r, ["Категория", "ИИ-помощник", "доля", "Аудит уязвимостей", "доля"], rows)
    wb_ = wb.create_sheet("Плохие запросы")
    XL.sheet_base(wb_, widths=[12, 20, 34, 80, 44, 9], title=title)
    r = XL.section(wb_, 1, "Запросы, которые не могут привести к результату", span=6)
    XL.table(wb_, r, ["Дата", "Раздел", "Категория", "Запрос", "Почему", "Повторов"],
             [[b["d"][:10], b["section"], QUALITY[b["quality"]], b["text"], b["why"] or "", b["n"]]
              for b in rep["bad_examples"]], wrap={"Запрос", "Почему"})
    wt = wb.create_sheet("Доработки")
    XL.sheet_base(wt, widths=[8, 12, 28, 80, 18], title=title)
    r = XL.section(wt, 1, "Предложения по доработкам (обращения)", span=5)
    XL.table(wt, r, ["№", "Дата", "Раздел", "Текст", "Статус"],
             [[t["id"], t["d"][:10], t["section"] or "", t["body"], t["status"]] for t in rep["tickets"]["items"]],
             wrap={"Текст"})
    wsrc = wb.create_sheet("Источники")
    XL.sheet_base(wsrc, widths=[8, 12, 26, 36, 22, 16], title=title)
    r = XL.section(wsrc, 1, "Предложения по новым источникам", span=6)
    XL.table(wsrc, r, ["№", "Дата", "Сайт", "Название", "Куда", "Статус"],
             [[x["id"], x["d"][:10], x["domain"], x["title"] or "", x["purpose"], x["status"]]
              for x in rep["sources"]["items"]], wrap={"Название"})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
