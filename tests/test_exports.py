"""Выгрузки в Excel и Word в стиле AuditLens: вкладка «Отзывы» и аудит-дело.

Без сети и БД: строки выгрузки и дело — фикстуры.
"""
from __future__ import annotations

import io
import re
import zipfile
from datetime import UTC, date, datetime

from openpyxl import load_workbook

from bank_audit.web import case_export, export_brand, reviews_export


def _rows(n: int = 40) -> list[dict]:
    themes = ["Блокировки по 161-ФЗ", "Оспаривание операций и возвраты (чарджбэк)",
              "Навязанные допуслуги (страховка, подписка)"]
    out = []
    for i in range(n):
        out.append({
            "дата": f"2026-0{7 + i % 3}-{1 + i % 27:02d}", "банк": "Сбербанк",
            "площадка": ["bankiru", "sravni_reviews", "finuslugi_reviews"][i % 3],
            "город": ["Москва", "Санкт-Петербург", ""][i % 3], "оценка": 1,
            "продукт": "Дебетовая карта", "главная проблема": themes[i % 3], "группа": "",
            "доп. проблемы": "", "эскалация": ["", "грозит", "обратился"][i % 3],
            "куда": "ЦБ" if i % 3 else "", "без согласия": "да" if i % 5 == 0 else "",
            "ввели в заблуждение": "", "уязвимый клиент": "пенсионер" if i % 7 == 0 else "",
            "сумма": 72600.0 if i == 1 else "", "дата события": "", "вне кодификатора": "",
            "суть": "Клиент требует вернуть деньги.",
            "цитата": "=HYPERLINK(\"http://evil\")" if i == 0 else "- банк отказал",
            "ссылка": f"https://www.banki.ru/services/responses/bank/response/{i}",
            "текст": "Полный текст отзыва."})
    return out


def test_reviews_xlsx_branded_with_charts_and_stamp():
    body = reviews_export.to_xlsx(_rows(), {"bank": "Сбербанк", "days": 90,
                                            "theme": None, "product": None})
    wb = load_workbook(io.BytesIO(body))
    assert wb.sheetnames == ["Обзор", "Жалобы", "Сводки", "О выгрузке"]
    assert wb.properties.creator == "AuditLens"
    ov = wb["Обзор"]
    assert len(ov._charts) == 4 and len(ov._images) == 1        # графики и шапка
    cells = [str(c.value) for row in ov.iter_rows() for c in row if c.value]
    assert any("Экспортировано из AuditLens" in c for c in cells)
    for name in wb.sheetnames:                                  # штамп и колонтитулы — везде
        ws = wb[name]
        assert "Экспортировано из AuditLens" in ws.oddFooter.left.text
        assert ws.oddHeader.left.text == "AuditLens"
    ws = wb["Жалобы"]
    head = [c.value for c in ws[4]]
    assert head[0] == "Дата" and "Цитата клиента" in head
    first = [c for c in ws[5]]
    assert isinstance(first[0].value, (date, datetime))         # дата — датой
    q = first[head.index("Цитата клиента")]
    assert q.data_type == "s" and q.value.startswith("=HYPERLINK")   # не формула
    second = [c.value for c in ws[6]]
    assert second[head.index("Цитата клиента")] == "- банк отказал"   # без апострофа
    assert second[head.index("Сумма, ₽")] == 72600.0
    assert ws.freeze_panes == "B5"
    ss = wb["Сводки"]
    vals = [c.value for row in ss.iter_rows() for c in row]
    assert "Блокировки по 161-ФЗ" in vals and "banki.ru" in vals


def _case() -> dict:
    items = [{"kind": "review", "url": f"https://www.banki.ru/r/{i}", "title": "снимок",
              "note": "проверить регламент" if i == 0 else "", "added_by": "a1",
              "review": {"bank": "Сбербанк", "date": "2026-09-20", "city": "Москва",
                         "source": "banki.ru", "product": "Вклад", "issue_label": "Проценты",
                         "risk": "conduct", "summary": "Суть жалобы", "quote": "=цитата",
                         "esc": "filed" if i == 0 else None, "esc_to": ["cbr"],
                         "vulnerable": [], "no_consent": False, "misled": False,
                         "amount": 1000.0 if i == 0 else None}} for i in range(3)]
    items.append({"kind": "document", "url": "https://www.sberbank.ru/x", "title": "Тариф",
                  "note": "", "added_by": "a2", "bank_name": "Сбербанк",
                  "fetched_at": "2026-09-24"})
    return {"case_id": 5, "owner": "a1", "title": "Проценты по вкладам", "note": "Цель",
            "analysis": "## Главное\n- пункт один\n**важно**", "shared": True,
            "updated_at": datetime(2026, 9, 26, 12, 0, tzinfo=UTC), "items": items}


def test_case_xlsx_branded():
    wb = load_workbook(io.BytesIO(case_export.to_xlsx(_case())))
    assert wb.sheetnames == ["Дело", "Материалы", "Сводки"]
    ov = wb["Дело"]
    assert len(ov._images) == 1 and len(ov._charts) >= 1
    cells = [str(c.value) for row in ov.iter_rows() for c in row if c.value]
    assert any("Экспортировано из AuditLens" in c for c in cells)
    assert any(c.startswith("•  пункт один") for c in cells)
    ws = wb["Материалы"]
    head = [c.value for c in ws[4]]
    q = ws.cell(row=5, column=head.index("Цитата") + 1)
    assert q.data_type == "s" and q.value == "=цитата"


def test_case_docx_branded_with_embedded_fonts():
    body = case_export.to_docx(_case())
    z = zipfile.ZipFile(io.BytesIO(body))
    names = z.namelist()
    fonts = [n for n in names if n.startswith("word/fonts/") and n.endswith(".odttf")]
    assert len(fonts) == 4                        # Geist ×2, Source Serif 4, JetBrains Mono
    table = z.read("word/fontTable.xml").decode()
    assert 'w:name="Geist"' in table and "w:embedRegular" in table and "w:embedBold" in table
    assert "embedTrueTypeFonts" in z.read("word/settings.xml").decode()
    assert "odttf" in z.read("[Content_Types].xml").decode()
    # обфускация обратима: ключ из fontTable возвращает настоящий TrueType
    rels = z.read("word/_rels/fontTable.xml.rels").decode()
    m = re.search(r'<w:embedRegular r:id="(\w+)" w:fontKey="(\{[^}]+\})"', table)
    target = re.search(rf'Id="{m.group(1)}"[^>]*Target="([^"]+)"', rels).group(1)
    raw = export_brand.obfuscate(z.read("word/" + target), m.group(2))
    assert raw[:4] == b"\x00\x01\x00\x00"         # TrueType
    footer = "".join(z.read(n).decode() for n in names if n.startswith("word/footer"))
    assert "Экспортировано из AuditLens" in footer and "NUMPAGES" in footer
    doc = z.read("word/document.xml").decode()
    assert "Дело в цифрах" in doc and "Материалы дела" in doc and "КОММЕНТАРИЙ АУДИТОРА" in doc
    assert "Source Serif 4 SemiBold" in doc or "Source Serif 4 SemiBold" in \
        z.read("word/styles.xml").decode()
    assert any(n.startswith("word/media/") for n in names)      # знак и графики
    core = z.read("docProps/core.xml").decode()
    assert "AuditLens" in core
