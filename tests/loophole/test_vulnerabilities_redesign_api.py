"""Сервер новой вкладки «Уязвимости»: сводка, фильтр проверки, сортировка,
карточка записи, суть только для находок, находки исследования, Excel.

Реальные HTTP-запросы и SQLite, модель — подменённая: суть не должна вызывать
модель для «не подтверждено» и должна получать текст с замаскированными ПД.
"""
from __future__ import annotations

import io
import zipfile

from sqlalchemy import text

from bank_audit.loophole import repository as repo
from bank_audit.loophole import summary as summary_mod
from bank_audit.loophole.models import LoopholeRecord
from tests.loophole import test_record_verdict_authorization as access
from tests.loophole.test_preliminary_research_source_import import _create_import_schema

session = access.session
client = access.client
API = "/api/loophole"


def _record(session, title, *, kind, bank="sberbank", conf=0.5, manual=False, raw=""):
    record_id = repo.insert_record(LoopholeRecord(
        sha256=f"sha-{title}", title=title, snippet=f"Фрагмент: {title}", bank_slug=bank,
        is_loophole=kind != "not_confirmed", classification=kind, status="preliminary",
        verdict_confidence=conf, verdict_reason=f"Комментарий классификатора: {title}",
        verdict_model="manual" if manual else "gpt", raw_text=raw or None,
        content_status="full" if raw else "legacy",
    ), session=session)
    return record_id


def _base(session):
    access._access(session)
    _create_import_schema(session)
    ids = {
        "vuln_wait": _record(session, "Кэшбэк за переводы", kind="vulnerability", conf=0.91),
        "vuln_done": _record(session, "Льготный период", kind="vulnerability", conf=0.6,
                             manual=True, bank="vtb"),
        "fraud_wait": _record(session, "Подмена QR-кода", kind="fraud_scheme", conf=0.86),
        "none_1": _record(session, "Списание за неактивность", kind="not_confirmed", conf=0.3),
        "none_2": _record(session, "Страховка", kind="not_confirmed", conf=0.4, bank="vtb"),
    }
    session.commit()
    return ids


def test_summary_counts_whole_base_and_facets_of_current_slice(client, session):
    _base(session)
    response = client.get(f"{API}/catalog/summary", headers=access._HEADERS,
                          params={"classification": "confirmed"})
    assert response.status_code == 200
    data = response.json()
    totals = data["totals"]
    assert totals["total"] == 5
    assert totals["vulnerability"] == 2 and totals["fraud_scheme"] == 1
    assert totals["not_confirmed"] == 2
    # «Ждут проверки» — правило очереди: находка модели без ручного вердикта.
    assert totals["awaiting"] == 2
    assert totals["awaiting_vulnerability"] == 1 and totals["awaiting_fraud_scheme"] == 1
    facets = data["facets"]
    assert facets["types"] == {"vulnerability": 2, "fraud_scheme": 1, "not_confirmed": 2,
                               "confirmed": 3, "all": 5}
    assert facets["awaiting"] == 2
    assert {b["slug"]: b["count"] for b in facets["banks"]} == {"sberbank": 2, "vtb": 1}

    only_vtb = client.get(f"{API}/catalog/summary", headers=access._HEADERS,
                          params={"classification": "confirmed", "bank_slugs": "vtb"}).json()
    assert only_vtb["facets"]["types"]["not_confirmed"] == 1
    assert only_vtb["totals"]["total"] == 5            # карточки — по всей базе


def test_verification_filter_follows_queue_rule_and_sort_by_confidence(client, session):
    ids = _base(session)
    awaiting = client.get(f"{API}/catalog", headers=access._HEADERS, params={
        "classification": "confirmed", "verification_status": "awaiting", "sort": "conf",
    }).json()["records"]
    assert [r["record_id"] for r in awaiting] == [ids["vuln_wait"], ids["fraud_wait"]]
    assert all(r["awaiting"] and not r["reviewed"] for r in awaiting)
    reviewed = client.get(f"{API}/catalog", headers=access._HEADERS, params={
        "classification": "confirmed", "verification_status": "reviewed",
    }).json()["records"]
    assert [r["record_id"] for r in reviewed] == [ids["vuln_done"]]
    assert reviewed[0]["reviewed"] is True
    bad = client.get(f"{API}/catalog", headers=access._HEADERS, params={"sort": "random"})
    assert bad.status_code == 422


def test_record_card_has_history_fields(client, session):
    ids = _base(session)
    data = client.get(f"{API}/records/{ids['vuln_done']}/content",
                      headers=access._HEADERS).json()
    for key in ("title", "classification", "verdict_model", "classifier_verdict_reason",
                "summary", "awaiting", "reviewed", "provenance", "decisions", "raw_text"):
        assert key in data
    assert data["reviewed"] is True and data["awaiting"] is False
    assert data["decisions"] == [] and data["provenance"] is None


class _FakeLLM:
    def __init__(self, answer="Переводы между своими картами проходят как покупки."):
        self.calls = []
        self.answer = answer

    async def ainvoke(self, messages):
        self.calls.append(messages)

        class _Response:
            content = self.answer
        return _Response()


def test_summary_only_for_findings_and_cached(client, session, monkeypatch):
    ids = _base(session)
    fake = _FakeLLM()
    monkeypatch.setattr(summary_mod, "_default_llm", lambda: fake)
    none = client.post(f"{API}/records/{ids['none_1']}/summary", headers=access._HEADERS)
    assert none.json() == {"summary": None, "generated": False, "reason": "not_finding"}
    assert fake.calls == []                      # «не подтверждено» — без вызова модели

    first = client.post(f"{API}/records/{ids['vuln_wait']}/summary", headers=access._HEADERS)
    assert first.json()["generated"] is True
    assert first.json()["summary"].startswith("Переводы между своими картами")
    second = client.post(f"{API}/records/{ids['vuln_wait']}/summary", headers=access._HEADERS)
    assert second.json()["generated"] is False
    assert len(fake.calls) == 1                  # суть составляется один раз
    card = client.get(f"{API}/records/{ids['vuln_wait']}/content", headers=access._HEADERS)
    assert card.json()["summary"] == first.json()["summary"]


def test_summary_masks_personal_data(client, session, monkeypatch):
    access._access(session)
    _create_import_schema(session)
    record_id = _record(session, "Звонки службы безопасности", kind="fraud_scheme",
                        raw="Звонили с номера +7 916 123-45-67, представились банком.")
    session.commit()
    fake = _FakeLLM()
    monkeypatch.setattr(summary_mod, "_default_llm", lambda: fake)
    client.post(f"{API}/records/{record_id}/summary", headers=access._HEADERS)
    sent = str(fake.calls[0])
    assert "916 123-45-67" not in sent


def test_catalog_xlsx_exports_visible_records_in_auditlens_style(client, session):
    ids = _base(session)
    response = client.post(f"{API}/export/catalog.xlsx", headers=access._HEADERS,
                           json={"classification": "confirmed"})
    assert response.status_code == 200
    assert "AuditLens_uyazvimosti_" in response.headers["content-disposition"]
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Обзор", "Записи", "Сводки"]
    assert book.properties.creator == "AuditLens"
    cells = [str(c.value) for row in book["Записи"].iter_rows() for c in row if c.value]
    assert "Кэшбэк за переводы" in cells and "Подмена QR-кода" in cells
    assert "Страховка" not in cells              # «не подтверждено» вне фильтра
    assert any("Экспортировано из AuditLens" in c for c in cells)
    picked = client.post(f"{API}/export/catalog.xlsx", headers=access._HEADERS,
                         json={"record_ids": [ids["none_2"]]})
    names = [str(c.value) for row in load_workbook(io.BytesIO(picked.content))["Записи"]
             .iter_rows() for c in row if c.value]
    assert "Страховка" in names
    assert zipfile.is_zipfile(io.BytesIO(response.content))


def test_research_findings_follow_import_links(client, session):
    ids = _base(session)
    workspace_id = repo.create_workspace(access._USERNAME, "Кэшбэк", session=session)
    session.execute(text(
        "INSERT INTO loophole_preliminary_import "
        "(research_id, source_id, workspace_id, record_id, imported_by) "
        "VALUES (7, 11, :ws, :rid, :u)"
    ), {"ws": workspace_id, "rid": ids["vuln_wait"], "u": access._USERNAME})
    session.commit()
    findings = client.get(f"{API}/research/workspace/{workspace_id}/findings",
                          headers=access._HEADERS).json()["findings"]
    assert [f["record_id"] for f in findings] == [ids["vuln_wait"]]
    assert findings[0]["awaiting"] is True
    history = client.get(f"{API}/history/{workspace_id}", headers=access._HEADERS).json()
    assert [f["record_id"] for f in history["findings"]] == [ids["vuln_wait"]]
    foreign = repo.create_workspace("someone-else", "Чужое", session=session)
    session.commit()
    denied = client.get(f"{API}/research/workspace/{foreign}/findings", headers=access._HEADERS)
    assert denied.status_code in (403, 404)


def test_queue_shows_model_verdict_in_server_order_with_total(client, session):
    from bank_audit.loophole import authorization

    access._access(session, role=authorization.ROLE_CCKS_EXPERT)
    _create_import_schema(session)
    old = _record(session, "Старая схема", kind="fraud_scheme", conf=0.6)
    fresh = _record(session, "Свежая уязвимость", kind="vulnerability", conf=0.95)
    _record(session, "Решённая", kind="vulnerability", conf=0.99, manual=True)
    _record(session, "Не подтверждено", kind="not_confirmed", conf=0.2)
    session.execute(text("UPDATE loophole_record SET collected_at = '2026-09-01T10:00:00' "
                         "WHERE record_id = :id"), {"id": old})
    session.execute(text("UPDATE loophole_record SET collected_at = '2026-09-20T10:00:00' "
                         "WHERE record_id = :id"), {"id": fresh})
    session.commit()

    by_age = client.get(f"{API}/queue", headers=access._HEADERS).json()
    assert [r["record_id"] for r in by_age["records"]] == [old, fresh]
    assert by_age["total"] == 2 and by_age["count"] == 2
    # Карточка очереди показывает, что предложила модель.
    assert by_age["records"][0]["classification"] == "fraud_scheme"
    assert by_age["records"][1]["verdict_model"] == "gpt"
    by_conf = client.get(f"{API}/queue", headers=access._HEADERS, params={"sort": "conf"}).json()
    assert [r["record_id"] for r in by_conf["records"]] == [fresh, old]
    bad = client.get(f"{API}/queue", headers=access._HEADERS, params={"sort": "random"})
    assert bad.status_code == 422


def test_new_findings_card_matches_the_seven_day_period_filter(client, session):
    from datetime import date, timedelta

    ids = _base(session)
    today = date.today()
    for key, days in (("vuln_wait", 2), ("fraud_wait", 10), ("none_1", 1)):
        session.execute(text("UPDATE loophole_record SET published_at = :d WHERE record_id = :id"),
                        {"d": (today - timedelta(days=days)).isoformat() + "T12:00:00", "id": ids[key]})
    session.commit()
    totals = client.get(f"{API}/catalog/summary", headers=access._HEADERS).json()["totals"]
    # «Не подтверждено» в новые находки не входит; прошлая неделя — отдельно.
    assert totals["new_7d"] == 1 and totals["new_prev_7d"] == 1
    listed = client.get(f"{API}/catalog", headers=access._HEADERS, params={
        "classification": "confirmed", "period_from": (today - timedelta(days=7)).isoformat(),
    }).json()
    assert listed["total"] == totals["new_7d"]
